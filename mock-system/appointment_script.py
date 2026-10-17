
"""Selenium automation for the local synthetic appointment mock only."""

import csv
import hashlib
import json
import logging
import os
import re
from datetime import datetime, date, time as dt_time
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import urlopen

import pandas as pd
from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import Select, WebDriverWait


BASE = Path(__file__).resolve().parent
EXCEL = BASE / "patients_appointments.xlsx"
RESULTS = BASE / "booking_results.csv"
ERRORS = BASE / "appointment_errors.log"

URL = "http://127.0.0.1:5000"
WAIT = 30

INPUT_COLUMNS = [
    "ID type", "PatientName", "ScheduleDate", "StartTime",
    "EndTime", "IsVirtual", "CLINICCODE", "DoctorName"
]

OUTPUT_COLUMNS = [
    "Excel_Row", "Row_Fingerprint", "Booking_Status",
    "Stage", "Details", "Updated_At"
]

KNOWN_ERRORS = {
    "No records found",
    "Patient Not Found",
    "Patient File Closed",
    "This patient file is suspended",
    "Incomplete Mobile Number",
    "No Third Name",
    "Appointment Already Exists",
    "No available slot",
    "Schedule Not Generated",
    "Patient Name Mismatch",
    "Overbooking Required",
    "Invalid Visit Type",
    "Invalid Date or Time",
    "Invalid Time Range",
    "Required Fields Missing"
}

REVIEW_STATUSES = {
    "Save Status Uncertain",
    "Review Required",
    "In Progress"
}

logging.basicConfig(
    filename=ERRORS,
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    encoding="utf-8"
)


class BookingError(Exception):
    def __init__(self, stage, message):
        self.stage = stage
        super().__init__(message)


def clean(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def normalize_date(value):
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.strftime("%Y-%m-%d")

    value = clean(value)

    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(value, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass

    raise BookingError("EXCEL", "Invalid Date or Time")


def normalize_time(value):
    if isinstance(value, (datetime, dt_time, pd.Timestamp)):
        return value.strftime("%H:%M")

    value = clean(value)

    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(value, fmt).strftime("%H:%M")
        except ValueError:
            pass

    raise BookingError("EXCEL", "Invalid Date or Time")


def prepare(row):
    national_id = clean(row["ID type"]).removesuffix(".0")

    if not re.fullmatch(r"900000000[1-7]", national_id):
        raise BookingError(
            "EXCEL",
            "Only synthetic patient IDs are permitted"
        )

    is_virtual = clean(row["IsVirtual"])

    if is_virtual in ("0", "0.0"):
        visit_type = "Followup Visit"
    elif is_virtual in ("1", "1.0"):
        visit_type = "TeleConsult"
    else:
        raise BookingError("EXCEL", "Invalid Visit Type")

    data = {
        "national_id": national_id,
        "patient_name": clean(row["PatientName"]),
        "appointment_date": normalize_date(row["ScheduleDate"]),
        "start_time": normalize_time(row["StartTime"]),
        "end_time": normalize_time(row["EndTime"]),
        "clinic_code": clean(row["CLINICCODE"]).upper(),
        "doctor_name": clean(row["DoctorName"]),
        "visit_type": visit_type
    }

    if not all(data.values()):
        raise BookingError("EXCEL", "Required Fields Missing")

    start = datetime.strptime(data["start_time"], "%H:%M")
    end = datetime.strptime(data["end_time"], "%H:%M")

    if (end - start).total_seconds() != 900:
        raise BookingError("EXCEL", "Invalid Time Range")

    data["source_key"] = "|".join([
        data["national_id"],
        data["clinic_code"],
        data["doctor_name"],
        data["appointment_date"],
        data["start_time"],
        data["end_time"]
    ])

    return data


def row_fingerprint(row):
    values = [clean(row[column]) for column in INPUT_COLUMNS]

    payload = json.dumps(
        values, ensure_ascii=False, separators=(",", ":")
    )

    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def load_results():
    if not RESULTS.exists():
        return {}

    with RESULTS.open(
        "r", newline="", encoding="utf-8-sig"
    ) as file:
        reader = csv.DictReader(file)

        if not set(OUTPUT_COLUMNS).issubset(reader.fieldnames or []):
            raise RuntimeError(
                "Existing booking_results.csv has a different format. "
                "Keep a backup and use a new results file."
            )

        return {
            int(item["Excel_Row"]): item
            for item in reader
        }


def save_results(results):
    temporary = RESULTS.with_suffix(".csv.tmp")

    with temporary.open(
        "w", newline="", encoding="utf-8-sig"
    ) as file:
        writer = csv.DictWriter(
            file, fieldnames=OUTPUT_COLUMNS
        )
        writer.writeheader()

        for row_number in sorted(results):
            writer.writerow(results[row_number])

        file.flush()
        os.fsync(file.fileno())

    os.replace(temporary, RESULTS)


EXCEL_RESULT_COLUMNS = ("Status", "Error_Reason", "Stage", "Updated_At")


def excel_status(booking_status):
    if booking_status.startswith("Successfully Saved"):
        return ("Successfully Registered", "")
    if booking_status in REVIEW_STATUSES:
        return ("Review Required", booking_status)
    return ("Failed", booking_status)


def sync_results_to_excel(results):
    """Mirror recorded CSV results to the source workbook, without altering inputs."""
    if not results:
        return

    workbook = load_workbook(EXCEL)
    sheet = workbook.active
    headers = {
        str(cell.value): cell.column
        for cell in sheet[1]
        if cell.value is not None
    }
    for name in EXCEL_RESULT_COLUMNS:
        if name not in headers:
            column = sheet.max_column + 1
            sheet.cell(row=1, column=column, value=name)
            headers[name] = column

    for row_number, result in results.items():
        if row_number < 2 or row_number > sheet.max_row:
            continue
        status, reason = excel_status(result["Booking_Status"])
        values = {
            "Status": status,
            "Error_Reason": reason if reason else "",
            "Stage": result["Stage"],
            "Updated_At": result["Updated_At"],
        }
        if result["Booking_Status"] in REVIEW_STATUSES:
            values["Error_Reason"] = result["Details"]
        elif status == "Failed":
            values["Error_Reason"] = result["Details"] or reason
        elif "OVER BOOKING" in result["Booking_Status"]:
            values["Status"] = "Successfully Registered - OVER BOOKING"
        for name, value in values.items():
            sheet.cell(row=row_number, column=headers[name], value=value)

    temporary = EXCEL.with_name(EXCEL.stem + ".tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, EXCEL)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()


def record(results, excel_row, fingerprint, status, stage, details):
    results[excel_row] = {
        "Excel_Row": excel_row,
        "Row_Fingerprint": fingerprint,
        "Booking_Status": status,
        "Stage": stage,
        "Details": details,
        "Updated_At": datetime.now().isoformat(
            timespec="seconds"
        )
    }

    save_results(results)
    sync_results_to_excel(results)


def api_get(endpoint, params=None):
    address = URL + endpoint

    if params:
        address += "?" + urlencode(params)

    with urlopen(address, timeout=WAIT) as response:
        return json.load(response)


def verify_booking(data):
    response = api_get(
        "/api/verify",
        {"source_key": data["source_key"]}
    )
    return bool(response.get("found"))


def check_slot(data):
    response = api_get(
        "/api/slot",
        {
            "clinic_code": data["clinic_code"],
            "doctor_name": data["doctor_name"],
            "appointment_date": data["appointment_date"],
            "start_time": data["start_time"]
        }
    )

    if not response.get("ok"):
        message = response.get("message") or "Unknown Error"
        raise BookingError("SLOT_CHECK", message)

    return response


def by_id(driver, name):
    return driver.find_element(By.ID, name)


def field_value(driver, name):
    return by_id(driver, name).get_attribute("value").strip()


def field_text(driver, name):
    return by_id(driver, name).text.strip()


def set_input(driver, name, value):
    field = by_id(driver, name)

    driver.execute_script(
        """
        const field = arguments[0];
        const value = arguments[1];
        const setter = Object.getOwnPropertyDescriptor(
            HTMLInputElement.prototype, "value"
        ).set;

        setter.call(field, value);

        field.dispatchEvent(
            new Event("input", {bubbles: true})
        );

        field.dispatchEvent(
            new Event("change", {bubbles: true})
        );
        """,
        field,
        value
    )

    if field_value(driver, name) != value:
        raise BookingError(
            "SET_FIELD",
            f"Unable to set {name}"
        )


def login(driver):
    driver.get(URL + "/booking")

    WebDriverWait(driver, WAIT).until(
        lambda d: by_id(d, "login_button").is_displayed()
    )

    set_input(driver, "login_username", "demo")
    set_input(driver, "login_password", "demo")

    by_id(driver, "login_button").click()

    WebDriverWait(driver, WAIT).until(
        lambda d: by_id(d, "dashboard").is_displayed()
    )


def choose_filters(driver, data):
    Select(
        by_id(driver, "clinic_code")
    ).select_by_value(data["clinic_code"])

    set_input(
        driver, "doctor_name", data["doctor_name"]
    )

    set_input(
        driver, "appointment_date", data["appointment_date"]
    )

    Select(
        by_id(driver, "visit_type")
    ).select_by_value(data["visit_type"])

    if field_value(driver, "clinic_code") != data["clinic_code"]:
        raise BookingError(
            "FILTERS", "Clinic selection mismatch"
        )

    if field_value(driver, "doctor_name") != data["doctor_name"]:
        raise BookingError(
            "FILTERS", "Care Provider Mismatch"
        )

    if field_value(driver, "appointment_date") != data["appointment_date"]:
        raise BookingError(
            "FILTERS", "Appointment Date Mismatch"
        )


def wait_for_calendar(driver, data):
    selector = (
        f'button.slot[data-date="{data["appointment_date"]}"]'
        f'[data-time="{data["start_time"]}"]'
    )

    def calendar_ready(browser):
        calendar = by_id(browser, "weekly_calendar")

        if calendar.get_attribute("data-ready") != "true":
            return False

        if field_value(browser, "clinic_code") != data["clinic_code"]:
            return False

        if field_value(browser, "doctor_name") != data["doctor_name"]:
            return False

        if field_value(browser, "appointment_date") != data["appointment_date"]:
            return False

        matches = browser.find_elements(
            By.CSS_SELECTOR, selector
        )

        return matches[0] if matches else False

    try:
        return WebDriverWait(
            driver, WAIT
        ).until(calendar_ready)

    except TimeoutException as exc:
        raise BookingError(
            "CALENDAR_LOAD",
            "Automation Timeout: calendar did not finish loading"
        ) from exc


def open_slot(driver, data):
    slot = wait_for_calendar(driver, data)

    slot_info = check_slot(data)
    existing_count = int(slot_info.get("count", 0))

    if not slot.is_enabled():
        raise BookingError(
            "SLOT_CHECK",
            "No available slot"
        )

    slot.click()

    WebDriverWait(driver, WAIT).until(
        lambda d: by_id(d, "booking_modal").is_displayed()
    )

    expected_fields = {
        "modal_date": data["appointment_date"],
        "start_time": data["start_time"],
        "end_time": data["end_time"]
    }

    for name, expected in expected_fields.items():
        if field_value(driver, name) != expected:
            raise BookingError(
                "MODAL_VERIFY",
                "Appointment Date/Time Mismatch"
            )

    if existing_count > 0:
        overbooking = WebDriverWait(driver, WAIT).until(
            lambda d: by_id(d, "overbook")
            if by_id(d, "overbook").is_displayed()
            else False
        )

        overbooking.click()

    return existing_count > 0


def lookup_patient(driver, data):
    set_input(
        driver, "national_id", data["national_id"]
    )

    def lookup_finished(browser):
        message = field_text(browser, "patient_message")
        name = field_value(browser, "patient_name")
        return message or name or False

    WebDriverWait(driver, WAIT).until(lookup_finished)

    message = field_text(driver, "patient_message")
    patient_name = field_value(driver, "patient_name")
    mobile = field_value(driver, "mobile")

    if message in KNOWN_ERRORS:
        raise BookingError("PATIENT_LOOKUP", message)

    if not patient_name:
        raise BookingError(
            "PATIENT_LOOKUP",
            message or "No records found"
        )

    if patient_name.casefold() != data["patient_name"].casefold():
        raise BookingError(
            "PATIENT_LOOKUP",
            "Patient Name Mismatch"
        )

    if len(patient_name.split()) < 3:
        raise BookingError("PRE_SAVE", "No Third Name")

    if not re.fullmatch(r"\d{10}", mobile):
        raise BookingError(
            "PRE_SAVE",
            "Incomplete Mobile Number"
        )

    if message != "Patient found.":
        raise BookingError(
            "PATIENT_LOOKUP",
            "Unknown Error: " + message
        )


def submit_booking(driver, data, overbooking):
    save_button = by_id(driver, "save")

    if not save_button.is_enabled():
        raise BookingError(
            "PRE_SAVE",
            "Save button is disabled"
        )

    save_button.click()

    try:
        response = WebDriverWait(driver, WAIT).until(
            lambda d: (
                field_text(d, "result")
                if field_text(d, "result")
                and field_text(d, "result")
                not in (
                    "Saving appointment...",
                    "Overbooking approved for this slot."
                )
                else False
            )
        )

    except TimeoutException as exc:
        raise BookingError(
            "POST_SAVE",
            "Save Status Uncertain: response timeout"
        ) from exc

    if response.startswith("Successfully Saved"):
        try:
            confirmed = verify_booking(data)
        except Exception as exc:
            raise BookingError(
                "VERIFY_SAVE",
                "Save Status Uncertain: verification failed"
            ) from exc

        if not confirmed:
            raise BookingError(
                "VERIFY_SAVE",
                "Save Status Uncertain: booking not found"
            )

        status = (
            "Successfully Saved — OVER BOOKING"
            if overbooking
            else "Successfully Saved"
        )

        return status, "VERIFY_SAVE", "Booking verified in mock database"

    if response in KNOWN_ERRORS:
        raise BookingError("POST_SAVE", response)

    raise BookingError(
        "POST_SAVE",
        "Unknown Error: " + response
    )


def process_booking(driver, data):
    try:
        if verify_booking(data):
            return (
                "Appointment Already Exists",
                "DUPLICATE_CHECK",
                "Existing booking found in mock database"
            )
    except Exception as exc:
        raise BookingError(
            "DUPLICATE_CHECK",
            "Unable to verify existing bookings"
        ) from exc

    login(driver)
    choose_filters(driver, data)

    overbooking = open_slot(driver, data)
    lookup_patient(driver, data)

    return submit_booking(driver, data, overbooking)

def classify_error(exc):
    stage = getattr(exc, "stage", "UNKNOWN")
    message = str(exc).strip() or type(exc).__name__

    if message in KNOWN_ERRORS:
        return message, stage, message

    if stage in ("POST_SAVE", "VERIFY_SAVE", "SAVE_CLICK"):
        return "Save Status Uncertain", stage, message

    if isinstance(exc, TimeoutException) or "Timeout" in message:
        return "Automation Timeout", stage, message

    if message.startswith("Unknown Error:"):
        return "Unknown Error", stage, message

    return "Automation Error", stage, message


def main():
    if not EXCEL.exists():
        raise FileNotFoundError(
            "patients_appointments.xlsx not found"
        )

    dataframe = pd.read_excel(
        EXCEL,
        dtype=str,
        keep_default_na=False
    )

    missing = [
        column for column in INPUT_COLUMNS
        if column not in dataframe.columns
    ]

    if missing:
        raise ValueError(
            "Missing Excel columns: " + ", ".join(missing)
        )

    results = load_results()
    sync_results_to_excel(results)  # Fill Excel with prior run results as well.

    driver = webdriver.Chrome()

    try:
        driver.set_page_load_timeout(WAIT)

        for index, row in dataframe.iterrows():
            excel_row = index + 2
            fingerprint = row_fingerprint(row)

            previous = results.get(excel_row)

            if previous:
                previous_hash = previous["Row_Fingerprint"]
                previous_status = previous["Booking_Status"]

                if previous_hash != fingerprint:
                    print(
                        f"Row {excel_row}: REVIEW REQUIRED - "
                        "Excel row changed; old result preserved"
                    )
                    continue

                if previous_status in REVIEW_STATUSES:
                    print(
                        f"Row {excel_row}: REVIEW REQUIRED - "
                        "check mock database before retrying"
                    )
                    continue

                print(
                    f"Row {excel_row}: SKIPPED - "
                    f"already processed ({previous_status})"
                )
                continue

            record(
                results,
                excel_row,
                fingerprint,
                "In Progress",
                "START",
                "Row processing started"
            )

            try:
                data = prepare(row)

                status, stage, details = process_booking(
                    driver, data
                )

            except Exception as exc:
                status, stage, details = classify_error(exc)

                logging.exception(
                    "Excel_Row=%s | Stage=%s | Status=%s | Details=%s",
                    excel_row, stage, status, details
                )

            record(
                results,
                excel_row,
                fingerprint,
                status,
                stage,
                details
            )

            print(
                f"Excel Row {excel_row}: "
                f"{status} | Stage: {stage}"
            )

    finally:
        driver.quit()


if __name__ == "__main__":
    main()
