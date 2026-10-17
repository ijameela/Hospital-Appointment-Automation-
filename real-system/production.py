"""RABET appointment automation: manual login and Excel-driven booking.

Developed by Jameelah Algethami.
Configure production with verified page locators before execution.
The script does not contain usernames, passwords, or patient records.
"""

import csv
import hashlib
import json
import logging
import os
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, time as dt_time
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import TimeoutException, WebDriverException
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select, WebDriverWait

# Resolve input and output files independently of the terminal location.
BASE = Path(__file__).resolve().parent
EXCEL = BASE / "patients_appointments.xlsx"
RESULTS = BASE / "booking_results.csv"
ERRORS = BASE / "appointment_errors.log"
CONFIG_FILE = BASE / "selectors.json"
WAIT = 30
MAX_PATIENTS_PER_SLOT = 3

INPUT_COLUMNS = [
    "ID type", "PatientName", "ScheduleDate", "StartTime",
    "EndTime", "IsVirtual", "CLINICCODE", "DoctorName",
]
OUTPUT_COLUMNS = [
    "Excel_Row", "Row_Fingerprint", "Booking_Status",
    "Stage", "Details", "Updated_At",
]
EXCEL_RESULT_COLUMNS = ("Status", "Error_Reason", "Stage", "Updated_At")
REVIEW_STATUSES = {"Save Status Uncertain", "Review Required", "In Progress"}
KNOWN_ERRORS = {
    "No records found", "Patient Not Found", "Patient File Closed",
    "This patient file is suspended", "Incomplete Mobile Number",
    "No Third Name", "Appointment Already Exists", "No available slot",
    "Schedule Not Generated", "Patient Name Mismatch",
    "Overbooking Limit Reached", "Invalid Visit Type",
    "Invalid Date or Time", "Invalid Time Range", "Required Fields Missing",
}

logging.basicConfig(
    filename=ERRORS, level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s", encoding="utf-8",
)


# Carry the exact workflow stage with each recognized failure.
class BookingError(Exception):
    def __init__(self, stage: str, message: str):
        self.stage = stage
        super().__init__(message)


# Store normalized appointment data in a single structure.
@dataclass(frozen=True)
class Appointment:
    national_id: str
    patient_name: str
    appointment_date: str
    start_time: str
    end_time: str
    clinic_code: str
    doctor_name: str
    visit_type: str

    @property
    def source_key(self) -> str:
        return "|".join((
            self.national_id, self.clinic_code, self.doctor_name,
            self.appointment_date, self.start_time, self.end_time,
        ))


# Read the actual page configuration without guessing selectors.
def load_config() -> dict[str, Any]:
    if not CONFIG_FILE.exists():
        raise SystemExit(
            "Missing production/selectors.json. Copy selectors.example.json "
            "to selectors.json and provide verified RABET locators."
        )
    with CONFIG_FILE.open(encoding="utf-8") as handle:
        config = json.load(handle)
    if not config.get("url") or "YOUR_" in config["url"]:
        raise SystemExit("Configure the real RABET login URL in selectors.json.")
    required = (
        "login_ready", "booking_page", "clinic", "doctor", "date", "visit_type",
        "calendar_ready", "slot", "slot_count", "slot_select", "overbook",
        "modal", "modal_date", "modal_start", "modal_end", "patient_id",
        "patient_name", "patient_mobile", "patient_message", "save", "result",
        "duplicate_search", "duplicate_count", "saved_booking_search", "saved_booking_count",
    )
    missing = [name for name in required if not config.get("selectors", {}).get(name)]
    if missing:
        raise SystemExit("Missing verified selectors: " + ", ".join(missing))
    if config.get("max_patients_per_slot") != MAX_PATIENTS_PER_SLOT:
        raise SystemExit("max_patients_per_slot must equal 3.")
    if not config.get("ready_url_fragment"):
        raise SystemExit("Configure ready_url_fragment to identify the authenticated page.")
    return config


# Normalize Excel values while preserving leading zeros in identifiers.
def clean(value: Any) -> str:
    return "" if pd.isna(value) else str(value).strip()


def normalize_date(value: Any) -> str:
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.strftime("%Y-%m-%d")
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(clean(value), fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise BookingError("EXCEL", "Invalid Date or Time")


def normalize_time(value: Any) -> str:
    if isinstance(value, (datetime, dt_time, pd.Timestamp)):
        return value.strftime("%H:%M")
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(clean(value), fmt).strftime("%H:%M")
        except ValueError:
            continue
    raise BookingError("EXCEL", "Invalid Date or Time")


# Validate appointment fields before opening the booking form.
def prepare(row: pd.Series) -> Appointment:
    national_id = clean(row["ID type"]).removesuffix(".0")
    if not re.fullmatch(r"\d{10}", national_id):
        raise BookingError("EXCEL", "Invalid patient ID")
    virtual = clean(row["IsVirtual"])
    if virtual in ("0", "0.0"):
        visit_type = "Followup Visit"
    elif virtual in ("1", "1.0"):
        visit_type = "TeleConsult"
    else:
        raise BookingError("EXCEL", "Invalid Visit Type")
    appointment = Appointment(
        national_id=national_id,
        patient_name=clean(row["PatientName"]),
        appointment_date=normalize_date(row["ScheduleDate"]),
        start_time=normalize_time(row["StartTime"]),
        end_time=normalize_time(row["EndTime"]),
        clinic_code=clean(row["CLINICCODE"]).upper(),
        doctor_name=clean(row["DoctorName"]),
        visit_type=visit_type,
    )
    if not all(vars(appointment).values()):
        raise BookingError("EXCEL", "Required Fields Missing")
    start = datetime.strptime(appointment.start_time, "%H:%M")
    end = datetime.strptime(appointment.end_time, "%H:%M")
    if (end - start).total_seconds() != 900:
        raise BookingError("EXCEL", "Invalid Time Range")
    return appointment


# Detect edits to previously processed Excel rows during Resume.
def row_fingerprint(row: pd.Series) -> str:
    payload = json.dumps([clean(row[col]) for col in INPUT_COLUMNS], ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


# Load prior outcomes so completed rows are never blindly repeated.
def load_results() -> dict[int, dict[str, str]]:
    if not RESULTS.exists():
        return {}
    with RESULTS.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if not set(OUTPUT_COLUMNS).issubset(reader.fieldnames or []):
            raise RuntimeError("Existing CSV has a different format; preserve it and review.")
        return {int(item["Excel_Row"]): item for item in reader}


# Persist the CSV atomically after every processed row.
def save_results(results: dict[int, dict[str, str]]) -> None:
    temp = RESULTS.with_suffix(".csv.tmp")
    with temp.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        for number in sorted(results):
            writer.writerow(results[number])
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temp, RESULTS)


# Mirror results into Excel without changing the appointment inputs.
def sync_results_to_excel(results: dict[int, dict[str, str]]) -> None:
    if not results:
        return
    workbook = load_workbook(EXCEL)
    try:
        sheet = workbook.active
        headers = {str(c.value): c.column for c in sheet[1] if c.value is not None}
        for name in EXCEL_RESULT_COLUMNS:
            if name not in headers:
                column = sheet.max_column + 1
                sheet.cell(1, column, name)
                headers[name] = column
        for row_number, result in results.items():
            if not 2 <= row_number <= sheet.max_row:
                continue
            status = result["Booking_Status"]
            if status.startswith("Successfully Saved"):
                display = "Successfully Registered - OVER BOOKING" if "OVER BOOKING" in status else "Successfully Registered"
                reason = result["Details"]
            elif status in REVIEW_STATUSES:
                display, reason = "Review Required", result["Details"]
            else:
                display, reason = "Failed", result["Details"] or status
            values = {
                "Status": display, "Error_Reason": reason,
                "Stage": result["Stage"], "Updated_At": result["Updated_At"],
            }
            for name, value in values.items():
                sheet.cell(row_number, headers[name], value)
        temp = EXCEL.with_name(EXCEL.stem + ".tmp.xlsx")
        try:
            workbook.save(temp)
            os.replace(temp, EXCEL)
        finally:
            if temp.exists():
                temp.unlink()
    finally:
        workbook.close()


# Record progress before and after each booking attempt.
def record(results: dict, excel_row: int, fingerprint: str, status: str, stage: str, details: str) -> None:
    results[excel_row] = {
        "Excel_Row": excel_row, "Row_Fingerprint": fingerprint,
        "Booking_Status": status, "Stage": stage, "Details": details,
        "Updated_At": datetime.now().isoformat(timespec="seconds"),
    }
    save_results(results)
    sync_results_to_excel(results)


# Resolve verified Selenium selectors from the local configuration.
class RabetUI:
    def __init__(self, driver: webdriver.Chrome, config: dict[str, Any]):
        self.driver = driver
        self.config = config
        self.wait = WebDriverWait(driver, WAIT)

    def locator(self, name: str) -> tuple[str, str]:
        spec = self.config["selectors"][name]
        method = spec.get("by", "css")
        choices = {"css": By.CSS_SELECTOR, "id": By.ID, "xpath": By.XPATH, "name": By.NAME}
        if method not in choices or not spec.get("value") or "YOUR_" in spec["value"]:
            raise BookingError("CONFIG", f"Invalid or placeholder selector: {name}")
        return choices[method], spec["value"]

    def element(self, name: str):
        return self.wait.until(EC.presence_of_element_located(self.locator(name)))

    def click(self, name: str) -> None:
        self.wait.until(EC.element_to_be_clickable(self.locator(name))).click()

    def text(self, name: str) -> str:
        return self.element(name).text.strip()

    def value(self, name: str) -> str:
        return (self.element(name).get_attribute("value") or "").strip()

    def fill(self, name: str, value: str) -> None:
        element = self.wait.until(EC.element_to_be_clickable(self.locator(name)))
        element.clear()
        element.send_keys(value)
        if self.value(name) != value:
            raise BookingError("SET_FIELD", f"Unable to set {name}")

    def select(self, name: str, value: str) -> None:
        element = self.element(name)
        if element.tag_name.lower() == "select":
            Select(element).select_by_value(value)
        else:
            raise BookingError("FILTERS", f"Selector {name} is not an HTML select; adapt this control")

    def wait_for_login(self) -> None:
        self.driver.get(self.config["url"])
        print("Sign in manually with your authorized username and password.")
        print("Waiting for the authenticated RABET page; no credentials are stored.")
        try:
            WebDriverWait(self.driver, 300).until(
                lambda d: self.config["ready_url_fragment"] in d.current_url
            )
            WebDriverWait(self.driver, 300).until(
                EC.visibility_of_element_located(self.locator("login_ready"))
            )
        except TimeoutException as exc:
            raise BookingError("LOGIN", "Manual Login Timeout") from exc

    def open_booking_page(self) -> None:
        self.click("booking_page")

    def choose_filters(self, data: Appointment) -> None:
        self.select("clinic", data.clinic_code)
        self.fill("doctor", data.doctor_name)
        self.fill("date", data.appointment_date)
        self.select("visit_type", data.visit_type)
        if self.value("clinic") != data.clinic_code:
            raise BookingError("FILTERS", "Clinic selection mismatch")
        if self.value("doctor") != data.doctor_name:
            raise BookingError("FILTERS", "Care Provider Mismatch")
        if self.value("date") != data.appointment_date:
            raise BookingError("FILTERS", "Appointment Date Mismatch")

    def query_count(self, action: str, count_field: str, data: Appointment, stage: str) -> int:
        # Search criteria must identify this exact patient/slot in the RABET UI.
        self.fill("patient_id", data.national_id) if action == "duplicate_search" else None
        self.click(action)
        try:
            count = int(self.text(count_field))
        except (ValueError, TimeoutException) as exc:
            raise BookingError(stage, "Unable to verify exact booking count") from exc
        if count < 0:
            raise BookingError(stage, "Invalid booking count")
        return count

    def duplicate_count(self, data: Appointment) -> int:
        # This action must search the same patient, clinic, doctor, date, and time.
        return self.query_count("duplicate_search", "duplicate_count", data, "DUPLICATE_CHECK")

    def slot_count(self, data: Appointment) -> int:
        # This element must reflect the live count for the exact selected slot.
        self.wait.until(EC.visibility_of_element_located(self.locator("calendar_ready")))
        try:
            count = int(self.text("slot_count"))
        except (ValueError, TimeoutException) as exc:
            raise BookingError("SLOT_CHECK", "Unable to determine live slot occupancy") from exc
        if count < 0:
            raise BookingError("SLOT_CHECK", "Invalid slot occupancy")
        return count

    def select_slot(self, data: Appointment, overbooking: bool) -> None:
        # Verified slot selector must be tied to the requested date and start time.
        slot = self.element("slot")
        expected_date = slot.get_attribute("data-date")
        expected_time = slot.get_attribute("data-time")
        if expected_date != data.appointment_date or expected_time != data.start_time:
            raise BookingError("SLOT_CHECK", "Selected slot does not match Excel date/time")
        self.click("slot_select")
        self.wait.until(EC.visibility_of_element_located(self.locator("modal")))
        if overbooking:
            self.click("overbook")
        if (self.value("modal_date") != data.appointment_date or
                self.value("modal_start") != data.start_time or
                self.value("modal_end") != data.end_time):
            raise BookingError("MODAL_VERIFY", "Appointment Date/Time Mismatch")

    def lookup_patient(self, data: Appointment) -> None:
        self.fill("patient_id", data.national_id)
        self.wait.until(lambda _: self.text("patient_message") or self.value("patient_name"))
        message = self.text("patient_message")
        name = self.value("patient_name")
        mobile = self.value("patient_mobile")
        if message in KNOWN_ERRORS:
            raise BookingError("PATIENT_LOOKUP", message)
        if not name:
            raise BookingError("PATIENT_LOOKUP", message or "No records found")
        if name.casefold() != data.patient_name.casefold():
            raise BookingError("PATIENT_LOOKUP", "Patient Name Mismatch")
        if len(name.split()) < 3:
            raise BookingError("PRE_SAVE", "No Third Name")
        if not re.fullmatch(r"\d{10}", mobile):
            raise BookingError("PRE_SAVE", "Incomplete Mobile Number")
        if message and message not in self.config.get("patient_success_messages", []):
            raise BookingError("PATIENT_LOOKUP", "Unknown Error: " + message)

    def verify_saved(self, data: Appointment) -> bool:
        # The configured search must verify the exact appointment, not just a toast.
        return self.query_count("saved_booking_search", "saved_booking_count", data, "VERIFY_SAVE") > 0

    def save(self, data: Appointment, overbooking: bool) -> tuple[str, str, str]:
        # Recheck the exact slot immediately before clicking Save.
        latest_count = self.slot_count(data)
        if latest_count >= MAX_PATIENTS_PER_SLOT:
            raise BookingError("PRE_SAVE", "Overbooking Limit Reached")
        if (latest_count > 0) != overbooking:
            raise BookingError("PRE_SAVE", "Slot occupancy changed; review required")
        try:
            self.click("save")
        except TimeoutException as exc:
            raise BookingError("SAVE_CLICK", "Save Status Uncertain: click timed out") from exc
        try:
            self.wait.until(lambda _: self.text("result") not in ("", "Saving appointment..."))
            response = self.text("result")
        except TimeoutException as exc:
            raise BookingError("POST_SAVE", "Save Status Uncertain: response timeout") from exc
        if response in KNOWN_ERRORS:
            raise BookingError("POST_SAVE", response)
        if response not in self.config.get("save_success_messages", []):
            raise BookingError("POST_SAVE", "Save Status Uncertain: unexpected response: " + response)
        try:
            confirmed = self.verify_saved(data)
        except Exception as exc:
            raise BookingError("VERIFY_SAVE", "Save Status Uncertain: verification failed") from exc
        if not confirmed:
            raise BookingError("VERIFY_SAVE", "Save Status Uncertain: booking not found")
        if overbooking:
            return "Successfully Saved - OVER BOOKING", "VERIFY_SAVE", "Appointment saved as OVER BOOKING"
        return "Successfully Saved", "VERIFY_SAVE", "Booking verified in RABET"


# Apply duplicate checks, occupancy rules, and patient validation.
def process_booking(ui: RabetUI, data: Appointment) -> tuple[str, str, str]:
    ui.open_booking_page()
    ui.choose_filters(data)
    if ui.duplicate_count(data) > 0:
        return "Appointment Already Exists", "DUPLICATE_CHECK", "Existing booking found"
    count = ui.slot_count(data)
    if count >= MAX_PATIENTS_PER_SLOT:
        return "Overbooking Limit Reached", "SLOT_CHECK", "Three patients already booked in this slot"
    overbooking = count > 0
    ui.select_slot(data, overbooking)
    ui.lookup_patient(data)
    return ui.save(data, overbooking)


# Classify failures without losing uncertain post-save outcomes.
def classify_error(exc: Exception) -> tuple[str, str, str]:
    stage = getattr(exc, "stage", "UNKNOWN")
    message = str(exc).strip() or type(exc).__name__
    if message in KNOWN_ERRORS:
        return message, stage, message
    if stage in ("SAVE_CLICK", "POST_SAVE", "VERIFY_SAVE"):
        return "Save Status Uncertain", stage, message
    if "review required" in message.lower():
        return "Review Required", stage, message
    if isinstance(exc, TimeoutException) or "Timeout" in message:
        return "Automation Timeout", stage, message
    if message.startswith("Unknown Error:"):
        return "Unknown Error", stage, message
    return "Automation Error", stage, message


# Execute each unprocessed row and preserve the resume ledger.
def main() -> None:
    config = load_config()
    if not EXCEL.exists():
        raise SystemExit("Place patients_appointments.xlsx in the production folder.")
    dataframe = pd.read_excel(EXCEL, dtype=str, keep_default_na=False)
    missing = [name for name in INPUT_COLUMNS if name not in dataframe.columns]
    if missing:
        raise SystemExit("Missing Excel columns: " + ", ".join(missing))
    results = load_results()
    sync_results_to_excel(results)
    driver = webdriver.Chrome()
    driver.set_page_load_timeout(WAIT)
    ui = RabetUI(driver, config)
    try:
        ui.wait_for_login()
        for index, row in dataframe.iterrows():
            excel_row = index + 2
            fingerprint = row_fingerprint(row)
            previous = results.get(excel_row)
            if previous:
                if previous["Row_Fingerprint"] != fingerprint:
                    print(f"Row {excel_row}: REVIEW REQUIRED - input changed")
                elif previous["Booking_Status"] in REVIEW_STATUSES:
                    print(f"Row {excel_row}: REVIEW REQUIRED - check RABET before retry")
                else:
                    print(f"Row {excel_row}: SKIPPED - {previous['Booking_Status']}")
                continue
            record(results, excel_row, fingerprint, "In Progress", "START", "Row processing started")
            try:
                appointment = prepare(row)
                status, stage, details = process_booking(ui, appointment)
            except Exception as exc:
                status, stage, details = classify_error(exc)
                logging.exception(
                    "Excel_Row=%s | Stage=%s | Status=%s | Details=%s",
                    excel_row, stage, status, details,
                )
            record(results, excel_row, fingerprint, status, stage, details)
            print(f"Excel Row {excel_row}: {status} | Stage: {stage}")
    finally:
        driver.quit()


if __name__ == "__main__":
    main()
